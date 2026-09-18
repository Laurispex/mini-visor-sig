from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).parent
OUT = ROOT / "Modelo_ER_Mini_Visor_SIG.docx"
ACTUAL = ROOT / "modelo_er_actual.png"
CONCEPTUAL = ROOT / "modelo_er_conceptual.png"

NAVY = "183B56"
TEAL = "087F72"
BLUE = "2E75B6"
MUTED = "5B6B7A"
LIGHT = "F4F7F9"
LINE = "AABBC8"
P_NAVY = f"#{NAVY}"
P_TEAL = f"#{TEAL}"
P_BLUE = f"#{BLUE}"
P_MUTED = f"#{MUTED}"
P_LINE = f"#{LINE}"


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_entity(draw, box, title, rows, fill):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=14, fill="#FFFFFF", outline=P_LINE, width=2)
    draw.rounded_rectangle((x1, y1, x2, y1 + 48), radius=14, fill=fill)
    draw.rectangle((x1, y1 + 30, x2, y1 + 48), fill=fill)
    draw.text((x1 + 18, y1 + 12), title, font=font(22, True), fill="#FFFFFF")
    y = y1 + 67
    for label, value, emphasis in rows:
        f = font(17, emphasis)
        draw.text((x1 + 18, y), label, font=f, fill=P_NAVY)
        draw.text((x1 + 150, y), value, font=font(17), fill=P_MUTED)
        y += 29


def relation(draw, a, b, label, vertical=False):
    if vertical:
        x = (a[0] + a[2]) // 2
        y1, y2 = a[3], b[1]
        draw.line((x, y1, x, y2), fill=P_TEAL, width=4)
        draw.ellipse((x - 5, y1 - 5, x + 5, y1 + 5), fill=P_TEAL)
        draw.ellipse((x - 5, y2 - 5, x + 5, y2 + 5), fill=P_TEAL)
        draw.text((x + 14, (y1 + y2) // 2 - 12), label, font=font(17, True), fill=P_TEAL)
    else:
        y = (a[1] + a[3]) // 2
        x1, x2 = a[2], b[0]
        draw.line((x1, y, x2, y), fill=P_TEAL, width=4)
        draw.ellipse((x1 - 5, y - 5, x1 + 5, y + 5), fill=P_TEAL)
        draw.ellipse((x2 - 5, y - 5, x2 + 5, y + 5), fill=P_TEAL)
        draw.text(((x1 + x2) // 2 - 34, y - 31), label, font=font(17, True), fill=P_TEAL)


def create_diagrams():
    # Modelo fisico actual: no se dibujan relaciones porque no existen FK en el esquema.
    img = Image.new("RGB", (1500, 720), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    draw.text((50, 35), "Modelo físico actual en Supabase PostGIS", font=font(30, True), fill=P_NAVY)
    draw.text((50, 82), "Las capas se almacenan como entidades espaciales independientes. No hay llaves foráneas definidas.", font=font(18), fill=P_MUTED)
    draw_entity(draw, (55, 170, 475, 490), "Capitales departamentales", [
        ("id", "integer", True), ("nombre", "varchar", False), ("departamento", "varchar", False), ("geom", "POINT SRID 4326", False),
    ], P_TEAL)
    draw_entity(draw, (540, 170, 960, 520), "Ciudades", [
        ("id", "integer", True), ("nombre", "varchar", False), ("departamento", "varchar", False), ("poblacion", "varchar", False), ("geom", "POINT SRID 4326", False),
    ], P_BLUE)
    draw_entity(draw, (1025, 150, 1445, 555), "Municipios atlas", [
        ("id", "integer", True), ("fid", "bigint", False), ("codigo", "varchar", False), ("codigo_departamento", "varchar", False), ("nombre", "varchar", False), ("poblacion_2018", "bigint", False), ("geom", "MULTIPOLYGON SRID 4326", False),
    ], P_NAVY)
    draw.rounded_rectangle((55, 610, 1445, 675), radius=10, fill="#E8F1F5", outline="#C7DBE4")
    draw.text((78, 631), "Nota: los campos departamento y codigo_departamento expresan una relación semántica, no una relación referencial en la base de datos.", font=font(16), fill=P_NAVY)
    img.save(ACTUAL)

    img = Image.new("RGB", (1500, 900), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    draw.text((50, 35), "Modelo conceptual propuesto", font=font(30, True), fill=P_NAVY)
    draw.text((50, 82), "Normaliza el departamento como entidad y formaliza las relaciones territoriales.", font=font(18), fill=P_MUTED)
    dep = (535, 160, 965, 385)
    mun = (70, 545, 500, 825)
    cap = (535, 545, 965, 795)
    city = (1000, 545, 1430, 825)
    draw_entity(draw, dep, "Departamentos", [
        ("id_departamento", "PK", True), ("codigo_departamento", "UNIQUE", False), ("nombre", "varchar", False), ("geom", "MULTIPOLYGON SRID 4326", False),
    ], P_TEAL)
    draw_entity(draw, mun, "Municipios", [
        ("id_municipio", "PK", True), ("codigo", "UNIQUE", False), ("id_departamento", "FK", False), ("nombre", "varchar", False), ("poblacion_2018", "bigint", False), ("geom", "MULTIPOLYGON SRID 4326", False),
    ], P_NAVY)
    draw_entity(draw, cap, "Capitales departamentales", [
        ("id_capital", "PK", True), ("id_departamento", "FK UNIQUE", False), ("nombre", "varchar", False), ("geom", "POINT SRID 4326", False),
    ], P_BLUE)
    draw_entity(draw, city, "Ciudades", [
        ("id_ciudad", "PK", True), ("id_departamento", "FK", False), ("nombre", "varchar", False), ("poblacion", "varchar", False), ("geom", "POINT SRID 4326", False),
    ], "#5B7894")
    relation(draw, dep, cap, "1 a 1", vertical=True)
    # Lines to left/right entities
    draw.line((dep[0] + 65, dep[3], mun[2] - 80, mun[1]), fill=P_TEAL, width=4)
    draw.text((315, 475), "1 a N", font=font(17, True), fill=P_TEAL)
    draw.line((dep[2] - 65, dep[3], city[0] + 80, city[1]), fill=P_TEAL, width=4)
    draw.text((1100, 475), "1 a N", font=font(17, True), fill=P_TEAL)
    draw.text((65, 855), "PK: llave primaria     FK: llave foránea     Las geometrías se conservan como tipos PostGIS con SRID 4326.", font=font(16), fill=P_MUTED)
    img.save(CONCEPTUAL)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(p, text, bold=False, color="000000", size=10.5):
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    return run


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = ["Entidad", "Campo", "Tipo", "Función"]
    for cell, value in zip(table.rows[0].cells, header):
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, value, True, "FFFFFF", 9.5)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, row_values)):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "EEF4F7")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, i == 0, NAVY if i == 0 else "283A4B", 9)
    doc.add_paragraph()


def build_document():
    create_diagrams()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(title, "Modelo Entidad Relación del Mini Visor SIG", True, "000000", 23)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(sub, "Supabase PostgreSQL PostGIS y capas geográficas de Colombia", False, MUTED, 11)
    doc.add_paragraph()

    p = doc.add_paragraph()
    add_text(p, "Propósito. ", True, NAVY)
    add_text(p, "Este documento describe el modelo de datos del Mini Visor SIG a partir de las tablas realmente disponibles en Supabase. Incluye el modelo físico actual y una propuesta conceptual de normalización para sustentar las relaciones territoriales.")

    doc.add_heading("Modelo físico actual", level=1)
    p = doc.add_paragraph()
    add_text(p, "Conclusión principal. ", True, NAVY)
    add_text(p, "El visor cuenta con tres entidades espaciales independientes. Las relaciones por departamento se expresan mediante atributos de texto; no existe una tabla Departamentos ni llaves foráneas en el esquema actual.")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(ACTUAL), width=Inches(6.9))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(cap, "Figura 1. Modelo físico actual en Supabase PostGIS.", False, MUTED, 9)

    doc.add_heading("Diccionario de entidades actuales", level=2)
    add_table(doc, [
        ("Capitales departamentales", "id", "integer", "Identificador de la entidad"),
        ("Capitales departamentales", "nombre y departamento", "varchar", "Nombre de capital y departamento asociado"),
        ("Capitales departamentales", "geom", "geometry POINT 4326", "Ubicación puntual"),
        ("Ciudades", "id", "integer", "Identificador de la entidad"),
        ("Ciudades", "nombre departamento poblacion", "varchar", "Atributos descriptivos"),
        ("Ciudades", "geom", "geometry POINT 4326", "Ubicación puntual"),
        ("Municipios atlas", "codigo y codigo_departamento", "varchar", "Códigos de municipio y departamento"),
        ("Municipios atlas", "nombre y poblacion_2018", "varchar y bigint", "Datos descriptivos y de población"),
        ("Municipios atlas", "geom", "geometry MULTIPOLYGON 4326", "Límite territorial municipal"),
    ])

    doc.add_heading("Modelo conceptual propuesto", level=1)
    p = doc.add_paragraph()
    add_text(p, "Propuesta. ", True, NAVY)
    add_text(p, "Para normalizar la información, se incorpora Departamentos como entidad territorial. Esta propuesta no reemplaza el modelo implementado; muestra cómo formalizar las relaciones que hoy están representadas por campos de texto.")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(CONCEPTUAL), width=Inches(6.9))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(cap, "Figura 2. Modelo conceptual propuesto para normalizar relaciones territoriales.", False, MUTED, 9)

    doc.add_heading("Relaciones propuestas", level=2)
    for text in [
        "Un departamento puede contener muchos municipios; cada municipio pertenece a un único departamento.",
        "Un departamento tiene una capital departamental; cada capital corresponde a un único departamento.",
        "Un departamento puede incluir varias ciudades; cada ciudad pertenece a un único departamento.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        add_text(p, text)

    doc.add_heading("Uso en el visor", level=2)
    p = doc.add_paragraph()
    add_text(p, "Las geometrías permanecen como tipos espaciales PostGIS. El backend consulta estas geometrías mediante ST_AsGeoJSON para que Leaflet las muestre en el navegador. La migración actual conserva EPSG 4326 en las tres capas y no altera el diseño del Mini Visor SIG.")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Mini Visor SIG - Modelo Entidad Relación", False, MUTED, 8.5)
    doc.save(OUT)


if __name__ == "__main__":
    build_document()
